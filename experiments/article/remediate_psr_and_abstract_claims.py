import docx

SRC = '/home/valentin/code/active/nexus/experiments/article/full paper marine policy (updated charts+text).docx'
OUT = SRC

d = docx.Document(SRC)

def set_paragraph_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''

CHANGES = []

p4 = d.paragraphs[4]      # Abstract
p15 = d.paragraphs[15]    # Introduction
p122 = d.paragraphs[122]  # Conclusion (biological origins + PSR sentence)

old4, old15, old122 = (p.text for p in (p4, p15, p122))

# ---------- Abstract: remove both PSR sentences, align two other claims with their body caveats ----------
assert 'evaluates the quality of inventions through Patent Search Reports' in old4
assert 'Most patents rely on established prior art, signaling incremental innovation.' in old4
new4 = old4
new4 = new4.replace(
    "the research maps the spatial distribution of patents, identifies leading organizations and biological "
    "sources, analyzes international patent extensions, and evaluates the quality of inventions through "
    "Patent Search Reports.",
    "the research maps the spatial distribution of patents, examines applicant concentration and biological "
    "sources -- each subject to caveats on data reconciliation detailed in the body -- and analyzes "
    "international patent extensions. An earlier version of this study also reported Patent Search Report "
    "(PSR) prior-art category statistics; that analysis has no traceable source dataset anywhere in the "
    "project archive and has been removed throughout the manuscript rather than restated unverified."
)
new4 = new4.replace(
    "Results show patent activity dominated by the private company Pharma Mar S.A. and a preference for a "
    "narrow set of marine organisms; a regional (sub-national) concentration analysis was attempted but its "
    "underlying data could not be independently verified from the project's archive and is not presented as "
    "a finding here.",
    "Results show patent activity concentrated among a small number of applicants, with Pharma Mar S.A. "
    "leading per a legacy applicant count that could only be partially reconciled against the project's "
    "archived data (see body), and a reliance on a narrow set of marine organism categories, whose "
    "underlying per-compound counts likewise have no independently reproducible source and are presented as "
    "illustrative rather than audited. A regional (sub-national) concentration analysis was attempted but its "
    "underlying data could not be independently verified from the project's archive and is not presented as "
    "a finding here."
)
new4 = new4.replace(
    "International patenting strategies reveal a focus on high-income markets, reinforcing global asymmetries "
    "in access to marine genetic resources. Most patents rely on established prior art, signaling incremental "
    "innovation.",
    "International patenting strategies reveal a focus on high-income markets, reinforcing global asymmetries "
    "in access to marine genetic resources."
)
assert new4 != old4
set_paragraph_text(p4, new4)
CHANGES.append(('P4 (Abstract)', old4, new4))

# ---------- Introduction: remove the PSR claim ----------
assert 'By analysing Patent Search Reports and mapping legal classifications' in old15
new15 = old15.replace(
    "By analysing Patent Search Reports and mapping legal classifications, the study provides a detailed "
    "understanding of the innovation quality and timing in this niche field—an area rarely explored in prior "
    "research.",
    "By mapping applicant concentration, biological source preferences, and international patenting "
    "strategies, the study provides an empirical view of innovation activity in this niche field—an area "
    "rarely explored in prior research."
)
set_paragraph_text(p15, new15)
CHANGES.append(('P15 (Introduction)', old15, new15))

# ---------- Conclusion: remove the PSR/X/Y/P-category claim ----------
assert 'high presence of X and Y category documents in Patent Search Reports' in old122
new122 = old122.replace(
    " The data further show that most patented inventions are built on well-established prior art (as "
    "revealed by the high presence of X and Y category documents in Patent Search Reports), indicating "
    "incremental rather than radical innovation trends. Additionally, the temporal dynamics of P-category "
    "documents stress the importance of timing in the patenting process.",
    " An earlier version of this manuscript reported Patent Search Report (X/Y/P-category) statistics in "
    "this paragraph as evidence of incremental innovation; that claim has been removed, consistent with its "
    "removal from Results and Discussion, since no dataset supporting it exists anywhere in the project's "
    "archive (see experiments/article/LEGACY_AUDIT_DECISIONS.md and "
    "experiments/article/POST_109_MANUSCRIPT_AUDIT.md)."
)
assert new122 != old122
set_paragraph_text(p122, new122)
CHANGES.append(('P122 (Conclusion)', old122, new122))

d.save(OUT)

print("=== SAVED ===")
for name, old, new in CHANGES:
    print("\n---", name, "---")
    print("OLD:", old[:250])
    print("NEW:", new[:250])
