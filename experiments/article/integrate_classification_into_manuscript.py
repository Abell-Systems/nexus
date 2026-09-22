import docx
from docx.shared import Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = '/home/valentin/code/active/nexus/experiments/article/full paper marine policy (updated charts+text).docx'
OUT = SRC  # overwrite in place; git tracks history

d = docx.Document(SRC)

def set_paragraph_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''

CHANGES = []

# Grab all target paragraph objects up front, by their ORIGINAL indices, before any
# insertion shifts later indices. Paragraph objects reference the underlying XML
# element directly, so they stay valid regardless of index shifts afterward.
p4 = d.paragraphs[4]
p53 = d.paragraphs[53]
p69 = d.paragraphs[69]
p74 = d.paragraphs[74]
p80 = d.paragraphs[80]
p81 = d.paragraphs[81]
old_tbl = d.tables[1]

old4, old53, old69, old74, old80, old81 = (p.text for p in (p4, p53, p69, p74, p80, p81))

# ---------- 1. Abstract: multi-source methodology wording ----------
assert 'Using data from the Espacenet database, the research maps' in old4
new4 = old4.replace(
    "Using data from the Espacenet database, the research maps the spatial distribution of patents, "
    "identifies leading organizations and biological sources, analyzes international patent extensions, "
    "and evaluates the quality of inventions through Patent Search Reports.",
    "Using patent data from the Espacenet database, complemented by a corrected, evidence-based relevance "
    "classification of Spain-relevant patent families built from the Minesoft Origin database, the research "
    "maps the spatial distribution of patents, identifies leading organizations and biological sources, "
    "analyzes international patent extensions, and evaluates the quality of inventions through Patent Search Reports."
)
set_paragraph_text(p4, new4)
CHANGES.append(('P4 (abstract)', old4, new4))

# ---------- 2. Methodology: state 12/14 compound scope ----------
assert 'After gathering the data from the Espacenet database' in old53
new_methodology_para = (
    "A corrected, evidence-based relevance classification supplements this analysis for Table 2 and "
    "Figures 7–9 (Section 4.2): every Spain-relevant patent family returned by a family-level search of "
    "the Minesoft Origin database (2000–2025) for each compound was individually screened against its full "
    "claim and description text and classified as Directly Relevant, Indirectly Relevant, Incidental Mention, "
    "or Uncertain, superseding an earlier, narrower country-code-filtered search. This corrected analysis "
    "covers 12 of the 14 approved marine-derived drugs listed in Table 1; Disitamab Vedotin and Tisotumab "
    "vedotin-tftv are not covered and are excluded from Table 2 and Figures 7–9."
)
new_p = d.add_paragraph(new_methodology_para, style=p53.style)
p53._p.addnext(new_p._p)
CHANGES.append(('New paragraph after P53 (methodology scope)', None, new_methodology_para))

# ---------- 3. Table 2 intro paragraph: units ----------
assert 'Table 2 presents the compounds that generated patents in total number of patents' in old69
new69 = (
    "Table 2 presents, for each of the 12 compounds covered by the corrected Minesoft-based classification "
    "(see Methodology), the number of Spain-relevant patent families classified as Directly Relevant or "
    "Indirectly Relevant, superseding the earlier Espacenet-based patent counts for this table. Of the 872 "
    "unique Spain-relevant families identified across all 12 compounds, 157 were classified as substantively "
    "relevant (91 Directly Relevant, 66 Indirectly Relevant); the remainder were Incidental Mentions (712) or "
    "Uncertain (3). Table 2 reports 161 compound–family observations rather than 157 unique families because "
    "four families matched more than one compound's search and are counted once per compound they matched."
)
set_paragraph_text(p69, new69)
CHANGES.append(('P69 (Table 2 intro)', old69, new69))

# ---------- 4. Rebuild Table 2 (index 1) with 4 columns ----------
rows_data = [
    ('Cytarabine', 69, 54, 123),
    ('Eribulin mesylate', 8, 0, 8),
    ('Brentuximab vedotin', 5, 2, 7),
    ('Trabectedin', 3, 3, 6),
    ('Vidarabine', 1, 3, 4),
    ('Polatuzumab vedotin', 4, 0, 4),
    ('Enfortumab vedotin', 2, 2, 4),
    ('Plitidepsin', 0, 2, 2),
    ('Omega 3 acid ethyl esters', 2, 0, 2),
    ('Ziconotide', 1, 0, 1),
    ('Omega 3 carboxylic acid', 0, 0, 0),
    ('Lurbinectedin', 0, 0, 0),
]
new_tbl = d.add_table(rows=len(rows_data) + 1, cols=4)
new_tbl.style = old_tbl.style
hdr = new_tbl.rows[0].cells
for cell, text in zip(hdr, ['Compound', 'Directly Relevant', 'Indirectly Relevant', 'Total Relevant']):
    cell.text = text
    cell.paragraphs[0].runs[0].font.size = Pt(11)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for ri, (name, di, ind, tot) in enumerate(rows_data, start=1):
    cells = new_tbl.rows[ri].cells
    for cell, val, align in zip(cells, [name, str(di), str(ind), str(tot)],
                                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]):
        cell.text = val
        cell.paragraphs[0].runs[0].font.size = Pt(11)
        cell.paragraphs[0].alignment = align
widths = [Emu(2324100), Emu(1300000), Emu(1300000), Emu(1300000)]
for col, w in zip(new_tbl.columns, widths):
    for cell in col.cells:
        cell.width = w

old_tbl._tbl.addprevious(new_tbl._tbl)
old_tbl._tbl.getparent().remove(old_tbl._tbl)
CHANGES.append(('Table 2 (rebuilt, 4 columns)',
                 'old 2-col Compound/Number of patents table',
                 'new 4-col Compound/Directly Relevant/Indirectly Relevant/Total Relevant table'))

# ---------- 5. Figure 3 prose: remove "leading compound" claim ----------
assert 'The leading compound in total number of patents, Brentuximab vedotin' in old74
new74 = (
    "In order to analyze the different compounds that generated each of the patents, according to their "
    "marine organism, a radial hierarchical diagram was generated as shown in Figure 3. Brentuximab vedotin, "
    "Polatuzumab vedotin, and Enfortumab vedotin share the same marine organism of origin, Mollusk/"
    "Cyanobacterium; under the corrected relevance classification (Table 2), Cytarabine is the leading "
    "compound by number of Spain-relevant patent families (123), followed by Eribulin mesylate (8) and "
    "Brentuximab vedotin (7). "
)
set_paragraph_text(p74, new74)
CHANGES.append(('P74 (Figure 3 prose)', old74, new74))

# ---------- 6. P80: organism-group ranking (also Table-2-derived; found while integrating, not in original audit) ----------
assert 'third largest compounds in the total number of patents' in old80
new80 = (
    "The Sponge marine organism originates three compounds in the analysis (Cytarabine, Vidarabine, and "
    "Eribulin mesylate); under the corrected classification, this group now accounts for the largest number "
    "of Spain-relevant patent families (135) of the five organism groups, ahead of Mollusk/Cyanobacterium "
    "(15). The Tunicate marine organism generates three compounds (Trabectedin, Plitidepsin, and "
    "Lurbinectedin), together accounting for 8 relevant families. The Cone snail is responsible for the "
    "Ziconotide-related patents (1 relevant family), and the fish-derived Omega-3 acid ethyl esters and "
    "Omega-3 carboxylic acid compounds together account for 2 relevant families. "
)
set_paragraph_text(p80, new80)
CHANGES.append(('P80 (organism-group ranking, Figure 3 discussion)', old80, new80))

# ---------- 7. P81: Figure 4 prose, "largest compound" + rate context ----------
assert 'The largest compound, Brentuximab vedotin, is vastly used for' in old81
tail = old81.split('while Cytarabine is dedicated to Leukemia.', 1)[1]
new81 = (
    "Figure 4 presents the alluvial graph of the disease areas where the compounds are used. The three "
    "leading compounds by number of Spain-relevant patent families under the corrected classification are "
    "Cytarabine (123), Eribulin mesylate (8), and Brentuximab vedotin (7), all dedicated to cancer treatments. "
    "Cytarabine's large absolute count partly reflects its much larger screened universe (690 Spain-relevant "
    "families screened, a ≈17.8% relevance rate) compared to other compounds; Eribulin mesylate, by contrast, "
    "has a substantially higher relevance rate (8 of 18 screened, ≈44%), a distinction the raw family counts "
    "alone do not convey. Brentuximab vedotin is vastly used for Anaplastic large T-cell systemic malignant "
    "lymphoma and Hodgkin’s disease treatments, while Cytarabine is dedicated to Leukemia." + tail
)
set_paragraph_text(p81, new81)
CHANGES.append(('P81 (Figure 4 prose)', old81, new81))

d.save(OUT)

print("=== SAVED ===")
for name, old, new in CHANGES:
    print("\n---", name, "---")
    if old is not None:
        print("OLD:", old[:220])
    print("NEW:", new[:220])
