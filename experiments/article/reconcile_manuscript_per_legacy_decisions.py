import docx

SRC = '/home/valentin/code/active/nexus/experiments/article/full paper marine policy (updated charts+text).docx'
OUT = SRC

d = docx.Document(SRC)

def set_paragraph_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''

CHANGES = []

p53 = d.paragraphs[53]
p75 = d.paragraphs[75]
p87 = d.paragraphs[87]   # "Table 3 - Number of patents per Spanish region" caption
p88 = d.paragraphs[88]   # "Source: primary data."
p90 = d.paragraphs[90]   # "Table 5 - Top 10 Spanish companies..." caption
p93 = d.paragraphs[93]   # Madrid/Pharma Mar + PSR paragraph (only PSR half touched here)
p112 = d.paragraphs[112] # PSR paragraph in Discussion
tbl3 = d.tables[2]        # Spanish region table

old53, old75, old87, old88, old90, old93, old112 = (
    p.text for p in (p53, p75, p87, p88, p90, p93, p112)
)

# ---------- 1. Methodology: correct the Espacenet/Minesoft provenance claim ----------
assert 'After gathering the data from the Espacenet database' in old53
new53 = old53.replace(
    "After gathering the data from the Espacenet database, the data was inspected in order to guarantee "
    "its usability for the applications to generate the images.",
    "The dataset behind Figures 1, 2, 5, and 6 and Table 5 was assembled via the Minesoft Origin database "
    "(number_lookup and bulk_publications), seeded from a curated list of 247 Spanish patent publication "
    "numbers whose original compilation is attributed to Espacenet; family, jurisdiction, legal-status, and "
    "applicant data for the 227 resulting unique families were then enriched via Minesoft. This is "
    "methodologically a second, differently-seeded Minesoft pull, not an independent Espacenet extraction "
    "run in parallel with the Section 4.2 classification below. After gathering this dataset, it was "
    "inspected in order to guarantee its usability for the applications used to generate the images."
)
set_paragraph_text(p53, new53)
CHANGES.append(('P53 (methodology provenance)', old53, new53))

# ---------- 2. Figure 3 prose: caveat the per-compound counts (does not touch Madrid/region claims) ----------
assert 'Cytarabine is the leading' in old75
new75 = old75.rstrip() + (
    " The organism groupings shown in Figure 3 (Mollusk/Cyanobacterium, Sponge, Tunicate, Cone snail, Fish) "
    "trace to Table 1's static compound-to-organism mapping; the per-compound patent counts underlying the "
    "figure's original layout, however, have no independently reproducible source in the project archive "
    "and should be read as illustrative rather than independently audited. "
)
set_paragraph_text(p75, new75)
CHANGES.append(('P75 (Figure 3 prose)', old75, new75))

# ---------- 3. Retire Table 3 -- minimal, mechanical note only; does NOT resolve the Madrid thesis ----------
new87 = (
    "Table 3 (Spanish region breakdown) has been retired: no sub-national region-level data exists in the "
    "project's archived datasets to independently verify or regenerate it (see "
    "experiments/article/LEGACY_AUDIT_DECISIONS.md). This affects claims made elsewhere in this manuscript "
    "about regional concentration within Spain; those claims are addressed separately, not resolved by this "
    "note alone."
)
set_paragraph_text(p87, new87)
set_paragraph_text(p88, '')
tbl3._tbl.getparent().remove(tbl3._tbl)
CHANGES.append(('P87/88 + Table 3 (retired, mechanical note only)', old87 + ' | ' + old88, new87))

# ---------- 4. Caveat Table 5 / Figure 5 discrepancy (does not touch Madrid/region claims) ----------
new_table5_caveat = (
    "Figure 5 and Table 5 below report different counts for the same top-applicants analysis: Figure 5 is "
    "regenerated directly and reproducibly from the archived dataset, while Table 5's counts (below) reflect "
    "an applicant-name homogenization step (merging corporate-entity spelling variants) whose exact method is "
    "not preserved in the project archive. A partial check found 4 of Table 5's 10 entries match the archived "
    "data's un-normalized counts exactly, confirming a shared source, but the remainder (e.g. Merck Patent "
    "Gmbh, Bayer AG, Pharma Mar SA) cannot currently be reconciled to Table 5's reported values even with "
    "generous name-variant merging -- a genuine data-completeness gap rather than a presentation difference. "
    "Both are retained pending an explicitly-labeled reconstruction, per the resolution in "
    "experiments/article/LEGACY_AUDIT_DECISIONS.md."
)
new_p = d.add_paragraph(new_table5_caveat, style=p90.style)
p90._p.addprevious(new_p._p)
CHANGES.append(('New paragraph before P90 (Table 5 caveat)', None, new_table5_caveat))

# ---------- 5. P93: remove ONLY the PSR claim; the Madrid/region sentence is left completely untouched here ----------
assert 'Madrid stands out as the leading region' in old93
assert '72% of the patents' in old93
old93_madrid_sentence = old93.split('Looking at the Patent Search Reports')[0].rstrip()
new93 = old93_madrid_sentence + (
    " The Patent Search Report (PSR) category figures previously reported in the remainder of this paragraph "
    "(X/Y/P/E percentages) have no traceable source dataset anywhere in the project archive -- unlike Table 2, "
    "which was corrected against a newly classified dataset, no dataset exists to correct this claim against, "
    "so it is removed here rather than restated unverified."
)
set_paragraph_text(p93, new93)
CHANGES.append(('P93 (PSR half only; Madrid sentence left untouched)', old93, new93))

# ---------- 6. P112: mirrored PSR paragraph in Discussion (does not touch Madrid claims) ----------
assert 'Patent Search Reports (PSRs) further illuminate' in old112
new112 = (
    "An earlier version of this manuscript reported Patent Search Report (PSR) category statistics (X/Y/P/E) "
    "in this section. That claim has been removed: no dataset supporting it exists anywhere in the project's "
    "archive (see experiments/article/LEGACY_AUDIT_DECISIONS.md), and it is not treated as correctable against "
    "a newer source the way Table 2 was, since no PSR dataset -- old or new -- is available to correct it "
    "against."
)
set_paragraph_text(p112, new112)
CHANGES.append(('P112 (Discussion PSR paragraph)', old112, new112))

d.save(OUT)

print("=== SAVED ===")
for name, old, new in CHANGES:
    print("\n---", name, "---")
    if old is not None:
        print("OLD:", old[:200])
    print("NEW:", new[:200])
