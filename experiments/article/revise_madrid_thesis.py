import docx

SRC = '/home/valentin/code/active/nexus/experiments/article/full paper marine policy (updated charts+text).docx'
OUT = SRC

d = docx.Document(SRC)

def set_paragraph_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''

CHANGES = []

p4 = d.paragraphs[4]     # Abstract
p94 = d.paragraphs[94]   # Results
p110 = d.paragraphs[110] # Discussion
p120 = d.paragraphs[120] # Conclusion

old4, old94, old110, old120 = (p.text for p in (p4, p94, p110, p120))

# ---------- Abstract ----------
assert 'Results show a high concentration of patent activity in the Community of Madrid' in old4
new4 = old4.replace(
    "Results show a high concentration of patent activity in the Community of Madrid, dominated by the "
    "private company Pharma Mar S.A., and a preference for a narrow set of marine organisms.",
    "Results show patent activity dominated by the private company Pharma Mar S.A. and a preference for a "
    "narrow set of marine organisms; a regional (sub-national) concentration analysis was attempted but its "
    "underlying data could not be independently verified from the project's archive and is not presented as "
    "a finding here."
)
set_paragraph_text(p4, new4)
CHANGES.append(('P4 (Abstract)', old4, new4))

# ---------- Results ----------
assert 'Madrid stands out as the leading region' in old94
new94 = old94.replace(
    "The results reveal that Madrid stands out as the leading region in Spain in terms of patent activity, "
    "with Pharma Mar S.A. as the main company in the country in terms of patent ownership.",
    "Pharma Mar S.A. appears as the leading company in Spain by patent ownership (Table 5; see the caveat "
    "above on Table 5's reconciliation with Figure 5). A regional breakdown by Spanish autonomous community "
    "was previously reported here (Table 3) but has been retired for lack of independently reproducible "
    "source data (see Methodology); no regional-concentration finding is asserted in its place."
)
set_paragraph_text(p94, new94)
CHANGES.append(('P94 (Results)', old94, new94))

# ---------- Discussion ----------
assert 'strongly concentrated in the Community of Madrid' in old110
new110 = (
    "At the national level, Spain's patent activity in this dataset is concentrated among a small number of "
    "institutions, with Pharma Mar S.A. as the leading applicant (Table 5; see the caveat on its "
    "reconciliation with Figure 5). This reflects the clustering of R&D infrastructure, financial capital, "
    "and institutional expertise in a small number of organizations. A sub-national, regional breakdown of "
    "this concentration (e.g. by autonomous community) was previously reported here but has been retired: no "
    "region-level data exists in the project's archive to independently verify or reconstruct it (see "
    "Methodology and experiments/article/LEGACY_AUDIT_DECISIONS.md). Whether this institutional concentration "
    "is also geographically concentrated within Spain -- and what that would imply for policies such as those "
    "suggested by Raimundo et al. (2018) to broaden scientific and technological capabilities across more "
    "autonomous communities -- is accordingly an open question this dataset cannot currently answer, not a "
    "finding of this study."
)
set_paragraph_text(p110, new110)
CHANGES.append(('P110 (Discussion)', old110, new110))

# ---------- Conclusion ----------
assert 'highly centralized in the Community of Madrid' in old120
new120 = (
    "Our findings highlight that marine pharmaceutical innovation in Spain is concentrated among a small "
    "number of institutions, largely driven by the activities of Pharma Mar S.A. This reflects strong "
    "institutional capacity and targeted investment. Whether this concentration also has a strong regional "
    "(sub-national) dimension within Spain is a question this study's archived data cannot currently answer "
    "-- an earlier regional breakdown (Table 3) was found to lack independently reproducible source data and "
    "has been retired; establishing this would require a dedicated, separately scoped regional dataset. "
    "Policy interventions to support knowledge diffusion and broader R&D capabilities remain relevant "
    "regardless, given the observed institutional concentration."
)
set_paragraph_text(p120, new120)
CHANGES.append(('P120 (Conclusion)', old120, new120))

d.save(OUT)

print("=== SAVED ===")
for name, old, new in CHANGES:
    print("\n---", name, "---")
    print("OLD:", old[:200])
    print("NEW:", new[:200])
